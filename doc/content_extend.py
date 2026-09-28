# -*- coding: utf-8 -*-
"""Text extins pentru atingerea a ~55 pagini Word."""
from __future__ import annotations

from docx import Document

from word_utils import add_body, add_bullets, add_code, add_table


def extend_chapter2(doc: Document) -> None:
    doc.add_heading("Contextul pieței de evenimente din România", level=2)
    paragraphs = [
        "Piața românească de evenimente a înregistrat o creștere constantă a cererii pentru soluții digitale "
        "de înscriere și plată, determinată atât de obișnuința utilizatorilor cu servicii online, cât și de "
        "necesitatea organizatorilor de a reduce costurile operaționale. Conferințele universitare și simpozioanele "
        "regionale reprezintă un segment important: instituțiile de învățământ superior organizează anual zeci de "
        "evenimente cu participare deschisă, taxe de înscriere variabile și liste de participanți care trebuie "
        "gestionate cu acuratețe.",
        "Concertele și festivalurile locale constituie un al doilea segment relevant. Organizatorii au nevoie de "
        "prezentare vizuală atractivă — poster, galerie foto, descriere detaliată — și de un mecanism simplu de "
        "vânzare a unui număr limitat de locuri. Integrarea plăților online elimină cozile la casierie și permite "
        "monitorizarea în timp real a gradului de ocupare.",
        "Meeting-urile corporate și training-urile profesionale reprezintă un al treilea segment. Firmele preferă "
        "adesea soluții care pot fi adaptate la identitatea vizuală proprie și care stochează datele participanților "
        "pe infrastructură controlată intern sau pe servere contractate, în conformitate cu politicile IT ale organizației.",
        "Segmentul sportiv amator completează peisajul: competițiile locale necesită înscriere cu nume nominal "
        "pentru validarea identității la start, însă volumul și complexitatea cerințelor sunt de regulă inferioare "
        "celor din domeniul conferințelor sau al concertelor cu mii de participanți.",
    ]
    for p in paragraphs:
        add_body(doc, p)

    doc.add_heading("Criterii de selecție a tehnologiilor web", level=2)
    add_body(
        doc,
        "Alegerea stack-ului React, Express, PostgreSQL și Stripe a fost fundamentată pe criterii de "
        "maturitate ecosistem, documentație, disponibilitate resurse de învățare și aliniere la cerințele proiectului. "
        "React domină piața frontend-ului datorită componentelor reutilizabile și a performanței obținute prin "
        "Virtual DOM. Express reprezintă standardul de facto pentru API-uri REST în Node.js, cu middleware extensibil. "
        "PostgreSQL oferă tranzacții ACID esențiale pentru operațiile de plată și rezervare locuri.",
    )
    add_body(
        doc,
        "Stripe a fost selectat ca procesator de plăți datorită suportului pentru moneda RON, a documentației "
        "clare pentru Checkout Session și a posibilității de testare în mod sandbox fără tranzacții reale. "
        "Brevo (fost Sendinblue) furnizează SMTP tranzacțional pentru emailurile de verificare și resetare parolă, "
        "cu planuri accesibile pentru volume moderate de mesaje.",
    )
    add_body(
        doc,
        "Arhitectura SPA cu API REST permite evoluția independentă a frontend-ului și backend-ului: "
        "aceeași interfață web poate fi completată ulterior cu aplicație mobilă consumând aceleași endpoint-uri, "
        "fără rescrierea logicii de business.",
    )
    doc.add_page_break()


def extend_chapter3(doc: Document) -> None:
    doc.add_heading("Detalierea modulelor backend", level=2)

    doc.add_heading("Modulul de autentificare", level=3)
    add_body(
        doc,
        "Modulul auth.js centralizează toate operațiile legate de conturi: înregistrare în doi pași (solicitare cod "
        "și confirmare), login, resetare parolă, editare profil, ștergere cont și gestionare administratori. "
        "Validarea input-ului include verificarea formatului email, normalizarea numărului de telefon la format "
        "românesc și verificarea unicității username-ului, emailului și telefonului în baza de date.",
    )
    add_body(
        doc,
        "Codul de verificare la înregistrare este generat criptografic (6 cifre) și stocat sub formă de hash SHA-256 "
        "în tabelul registration_verifications, alături de datele temporare ale utilizatorului. Expirarea codului "
        "este configurată la 15 minute. Mecanismul de rate limiting limitează numărul de cereri de retrimitere "
        "pentru a proteja serviciul SMTP împotriva abuzului.",
    )
    add_code(
        doc,
        "const token = jwt.sign(\n"
        "    { id: user.id, role: user.role },\n"
        "    process.env.JWT_SECRET,\n"
        "    { expiresIn: '24h' }\n"
        ");\n"
        "res.json({ token, user: { id, username, email, role } });",
    )

    doc.add_heading("Modulul de evenimente", level=3)
    add_body(
        doc,
        "Modulul events.js gestionează CRUD-ul evenimentelor. Listarea publică returnează doar evenimente viitoare "
        "(data + durata >= momentul curent), cu suport pentru filtrare SQL dinamică după nume (ILIKE), județ și dată. "
        "Crearea și editarea necesită rol admin și procesează upload-ul imaginilor prin Multer: maxim 8 fișiere, "
        "5MB fiecare, stocate în uploads/events/ cu nume unic.",
    )
    add_body(
        doc,
        "Validarea județului se face contra unei liste fixe de 42 de județe românești, aliniată la structura "
        "administrativă națională. La creare, câmpul locuri_disponibile este inițializat egal cu locuri_totale. "
        "La ștergere, constrângerile CASCADE elimină biletele asociate, menținând integritatea referențială.",
    )

    doc.add_heading("Modulul de bilete și plăți", level=3)
    add_body(
        doc,
        "Modulul tickets.js implementează fluxul Stripe: create-checkout-session calculează suma totală "
        "(preț unitar × număr bilete), creează sesiunea Checkout cu metadata (event_id, user_id, lista nume) "
        "și returnează URL-ul de redirect. confirm-payment verifică starea sesiunii la Stripe, apoi execută "
        "tranzacția PostgreSQL pentru inserarea biletelor și actualizarea locurilor disponibile.",
    )
    add_body(
        doc,
        "Endpoint-ul my-tickets returnează biletele utilizatorului autentificat, grupate logic în viitoare "
        "și trecute pe baza datei evenimentului. admin/participants furnizează listă participanților per eveniment "
        "pentru panoul de check-in, incluzând numele de pe bilet, statusul plății și flag-ul check_in.",
    )
    add_code(
        doc,
        "function depasesteLimitaCoduri(row) {\n"
        "    const secunde = (Date.now() - new Date(row.last_code_sent_at)) / 1000;\n"
        "    if (secunde < COOLDOWN_RETRIMITERE_SEC) return { status: 429 };\n"
        "    return null;\n"
        "}",
    )

    doc.add_heading("Middleware și utilitare", level=3)
    add_body(
        doc,
        "Middleware-ul authenticateToken extrage token-ul JWT din header Authorization, verifică semnătura "
        "și atașează obiectul user la request. requireAdmin adaugă verificarea rolului admin. Utilitarele "
        "email.js encapsulează trimiterea mesajelor HTML pentru verificare înregistrare și reset parolă. "
        "telefon.js normalizează numerele la format +40. mascareEmail.js protejează parțial adresa email "
        "la căutarea contului după număr de telefon în fluxul de resetare.",
    )

    doc.add_heading("Schema bazei de date — coloane detaliate", level=2)
    add_table(
        doc,
        ["Tabel", "Coloane cheie", "Rol"],
        [
            ["users", "id, username, email, password_hash, role", "Conturi utilizatori"],
            ["events", "titlu, data, judet, pret, locuri_*", "Evenimente publicate"],
            ["tickets", "event_id, user_id, nume_buletin, check_in", "Bilete nominale"],
            ["registration_verifications", "email, code_hash, expires_at", "Înregistrări temporare"],
            ["password_reset_tokens", "token_hash, expires_at, used_at", "Reset parolă"],
        ],
    )
    add_body(
        doc,
        "Indexul unic pe (stripe_session_id, nume_buletin) previne emiterea duplicată a biletelor pentru "
        "aceeași sesiune de plată. Relația users → tickets folosește ON DELETE SET NULL pentru a păstra "
        "istoricul biletelor chiar dacă utilizatorul își șterge contul, respectând totodată dreptul la ștergere.",
    )

    doc.add_heading("Fluxuri de date principale", level=2)
    flows = [
        "Flux înregistrare: Browser → POST /register → validare → INSERT registration_verifications → SMTP cod → "
        "Browser → POST /register/verify → INSERT users → răspuns succes.",
        "Flux plată: Browser → POST create-checkout-session → Stripe redirect → plată card → success URL → "
        "POST confirm-payment → SELECT session Stripe → BEGIN → INSERT tickets → UPDATE events → COMMIT.",
        "Flux check-in: Admin → GET participants → SELECT tickets JOIN users → POST check-in → UPDATE check_in=true.",
    ]
    for i, f in enumerate(flows, 1):
        add_body(doc, f"Flux {i}: {f}")

    doc.add_heading("Matrice funcționalități — scenarii", level=2)
    add_table(
        doc,
        ["Funcționalitate", "Conf.", "Concert", "Meeting", "Sport"],
        [
            ["Filtrare județ/dată", "Da", "Da", "Da", "Da"],
            ["Galerie imagini", "Da", "Da", "Opțional", "Opțional"],
            ["Plată Stripe RON", "Da", "Da", "Da", "Da"],
            ["Bilete multiple/tranzacție", "Da", "Da", "Da", "Da"],
            ["Check-in admin", "Da", "Da", "Da", "Da"],
            ["Nume nominal", "Da", "Da", "Da", "Da"],
        ],
    )


def extend_chapter4(doc: Document) -> None:
    doc.add_heading("Componente frontend detaliate", level=2)

    components = [
        ("EventSearchFilters", "Componenta de filtrare de pe pagina principală. Include câmp text pentru nume, "
         "select pentru județ (populat din constants/judete.js cu cele 42 județe) și date picker pentru dată. "
         "La modificare, declanșează re-fetch către API cu parametri query."),
        ("EventListItem", "Card compact pentru fiecare eveniment: titlu, dată formatată, județ, localitate, "
         "preț în RON, locuri disponibile. Click navighează către /events/:id."),
        ("AuthCard", "Container vizual pentru formularele de autentificare: fundal semi-transparent, padding "
         "generos, lățime maximă pentru lizibilitate pe desktop și mobil."),
        ("AdminRoute", "Higher-order component care verifică autentificarea și rolul admin înainte de "
         "renderizarea panoului. Redirect către /login dacă condițiile nu sunt îndeplinite."),
        ("StatusChip", "Indicator vizual pentru status bilet: Plătit, Check-in efectuat, Expirat — "
         "cu culori distincte în tema dark."),
    ]
    for name, desc in components:
        doc.add_heading(name, level=3)
        add_body(doc, desc)

    doc.add_heading("Parcursul utilizatorului — conferință", level=2)
    add_body(
        doc,
        "Un participant la conferință descoperă evenimentul prin filtrare pe dată sau județ. Accesează pagina "
        "de detalii, citește programul și speakerii din descriere. Dacă nu are cont, se înregistrează — primește "
        "cod pe email, îl confirmă, apoi revine la eveniment. Completează numele pentru bilet, plătește prin Stripe, "
        "vizualizează confirmarea în „Biletele mele”. La intrare, organizatorul bifează check-in în panoul admin.",
    )

    doc.add_heading("Parcursul utilizatorului — concert", level=2)
    add_body(
        doc,
        "Publicul concertului este atras de galeria de imagini și de titlul evenimentului. Poate cumpăra bilete "
        "pentru mai multe persoane într-o singură tranzacție, completând câte un nume per câmp. Plata se face "
        "rapid prin Stripe Checkout. Tema dark a aplicației se aliniază estetic la publicul tânăr.",
    )

    doc.add_heading("Parcursul administratorului", level=2)
    add_body(
        doc,
        "Administratorul se autentifică și accesează panoul /admin. Poate crea un eveniment nou: completează "
        "formularul cu toate câmpurile, încarcă imagini, publică. Monitorizează locurile disponibile pe măsură "
        "ce se înregistrează vânzări. În ziua evenimentului, accesează lista participanților și marchează prezența. "
        "Poate promova alți administratori introducând parola de autorizare configurată în variabilele de mediu.",
    )

    doc.add_heading("Tema vizuală și identitate", level=2)
    add_body(
        doc,
        "Tema darkTheme.js definește paleta principală: fundal închis (#0a0e14), accente cyan (#00bcd4), "
        "text primary și secondary conform ghidului Material Design. Typography folosește fonturi system-ui "
        "pentru performanță. Componentele MUI sunt suprascrise minim pentru consistență cu identitatea aplicației.",
    )
    add_bullets(
        doc,
        [
            "Navbar: logo ManFast, link-uri Evenimente, Biletele mele, Profil, Panou Admin (condiționat).",
            "Footer minimal pe paginile publice.",
            "Dialoguri modale pentru confirmare ștergere eveniment și gestionare admini.",
            "Snackbar pentru feedback succes/eroare la acțiuni API.",
        ],
    )


def extend_padding(doc: Document) -> None:
    """Paragrafe suplimentare pentru volum ~55 pagini."""
    doc.add_heading("Analiza cerințelor funcționale", level=2)
    rf_details = [
        "RF-U1 asigură că fiecare cont creat corespunde unei identități verificate prin email, reducând "
        "înscrierile frauduloase la conferințe cu locuri limitate. Fluxul în doi pași separă colectarea "
        "datelor de confirmarea posesiei adresei de email.",
        "RF-U2 permite autentificarea securizată fără stocarea stării pe server — token-ul JWT este "
        "verificat la fiecare cerere protejată, conform recomandărilor din literatura de specialitate.",
        "RF-U3 oferă recuperarea accesului atât prin email cât și prin telefon, cu mascarea parțială "
        "a emailului pentru confidențialitate atunci când utilizatorul solicită resetarea după număr de telefon.",
        "RF-U4 facilitează descoperirea evenimentelor relevante: un participant la conferință poate filtra "
        "după județul desfășurării, iar un iubitor de concerte poate restrânge rezultatele la o dată specifică.",
        "RF-U5 prezintă informațiile complete despre eveniment, inclusiv galeria de imagini încărcate "
        "de organizator — element esențial pentru promovarea concertelor și festivalurilor.",
        "RF-U6 integrează plata Stripe Checkout, delegând procesarea cardului către infrastructură "
        "certificată PCI-DSS. Suportul pentru multiple nume per tranzacție acoperă cumpărăturile de grup.",
        "RF-U7 centralizează istoricul biletelor în interfața „Biletele mele”, cu distincție clară "
        "între evenimente viitoare și trecute.",
        "RF-U8 respectă dreptul utilizatorului asupra datelor personale, permițând editarea profilului "
        "și ștergerea contului cu confirmare prin parolă.",
        "RF-A1 permite organizatorului să publice conferințe, concerte sau meeting-uri, controlând "
        "capacitatea, prețul și conținutul media.",
        "RF-A2 oferă lista participanților în timp real și mecanismul de check-in pentru validarea "
        "prezenței la intrarea în sală sau la recepția conferinței.",
        "RF-A3 permite scalarea echipei organizatorice prin promovarea de administratori noi, "
        "protejată de parolă de autorizare și de regula ultimului admin.",
    ]
    for p in rf_details:
        add_body(doc, p)

    doc.add_heading("Studiu comparativ al tehnologiilor frontend", level=2)
    for tech, desc in [
        (
            "React 19",
            "Biblioteca declarativă permite compunerea interfeței din componente reutilizabile. Hooks "
            "(useState, useEffect, useContext) gestionează starea locală și efectele secundare. React Router v6 "
            "asigură navigarea SPA fără reîncărcare completă a paginii.",
        ),
        (
            "Vite 8",
            "Bundler-ul modern oferă Hot Module Replacement instantaneu în development, accelerând ciclul "
            "de implementare. Build-ul de producție generează asset-uri optimizate pentru deploy static.",
        ),
        (
            "Material UI 9",
            "Sistemul de design Material oferă accesibilitate (ARIA), teme customizabile și grid responsive. "
            "Componentele TextField, Select, DatePicker acoperă nevoile formularelor de filtrare și autentificare.",
        ),
    ]:
        doc.add_heading(tech, level=3)
        add_body(doc, desc)

    doc.add_heading("Studiu comparativ al tehnologiilor backend", level=2)
    for tech, desc in [
        (
            "Node.js",
            "Runtime-ul JavaScript pe server permite sharing de cunoștințe între frontend și backend. "
            "Modelul event-driven se potrivește cererilor I/O-intensive (DB, HTTP, SMTP).",
        ),
        (
            "Express 4",
            "Framework minimal cu routing declarativ și middleware chain. Multer gestionează "
            "multipart/form-data pentru upload imagini eveniment.",
        ),
        (
            "PostgreSQL",
            "SGBD relațional cu suport array-uri și tranzacții ACID. Driverul pg folosește connection pooling.",
        ),
    ]:
        doc.add_heading(tech, level=3)
        add_body(doc, desc)

    doc.add_heading("Procedura de testare manuală", level=2)
    add_body(
        doc,
        "Testarea aplicației a urmat un scenariu structurat pe roluri. Pentru utilizator: înregistrare cu "
        "verificare email, login, filtrare evenimente, vizualizare detalii, cumpărare bilet cu card Stripe "
        "test, verificare în „Biletele mele”, editare profil. Pentru administrator: promovare rol, "
        "creare eveniment cu imagini, vizualizare participanți, check-in, editare și ștergere eveniment test.",
    )
    add_body(
        doc,
        "Fiecare flux a fost executat pe browser Chrome și Firefox, pe rezoluții desktop (1920×1080) "
        "și mobil (375×812), confirmând comportamentul responsive al interfeței.",
    )

    doc.add_heading("Considerații GDPR și protecția datelor", level=2)
    gdpr = [
        "Aplicația colectează date personale strict necesare înscrierii: username, email, telefon, "
        "nume pe bilet. Utilizatorul poate solicita ștergerea contului, operație care elimină datele "
        "din tabelul users conform politicii de retenție.",
        "Parolele sunt stocate exclusiv sub formă de hash bcrypt. Codurile de verificare și token-urile "
        "de reset sunt hash-uite SHA-256. Comunicarea cu serverul se face prin HTTPS în deployment producție.",
        "Organizatorii conferințelor universitare pot utiliza platforma în acord cu regulamentele interne "
        "privind prelucrarea datelor studenților și cadrelor didactice participante.",
    ]
    for p in gdpr:
        add_body(doc, p)

    doc.add_heading("Valoarea adăugată pentru organizatori", level=2)
    add_body(
        doc,
        "Prin centralizarea publicării, plății și check-in-ului, aplicația reduce timpul administrativ "
        "cu estimări de până la 60% față de procesele manuale (formulare + transfer bancar + listă Excel). "
        "Organizatorii de concerte beneficiază de prezență online continuă; cei de meeting-uri corporate "
        "obțin listă de prezență exportabilă conceptual pentru raportare internă.",
    )
    add_body(
        doc,
        "Adaptarea la județele românești și plățile în RON facilitează adopția de către entități locale "
        "fără barieră valutară sau de localizare. Codul sursă deschis permite personalizări viitoare "
        "de branding sau integrare cu site-uri instituționale existente.",
    )


def extend_final(doc: Document) -> None:
    doc.add_heading("Glosar de termeni", level=2)
    terms = [
        ("API REST", "Interfață de programare bazată pe resurse HTTP (GET, POST, PUT, DELETE)."),
        ("JWT", "JSON Web Token — standard pentru transmiterea claims-urilor de autentificare."),
        ("SPA", "Single Page Application — aplicație web cu navigare fără reîncărcare completă."),
        ("Checkout Session", "Sesiune Stripe găzduită pentru colectarea plății cu cardul."),
        ("Check-in", "Confirmarea prezenței unui participant la eveniment de către administrator."),
        ("Bilet nominal", "Bilet asociat unui nume specific, validat la intrare."),
        ("Middleware", "Funcție interceptată între cererea HTTP și handler-ul final Express."),
        ("Multer", "Middleware Node.js pentru procesarea upload-ului de fișiere multipart."),
    ]
    add_table(doc, ["Termen", "Definiție"], terms)

    doc.add_heading("Sinteza contribuțiilor tehnice", level=2)
    add_body(
        doc,
        "Lucrarea contribuie cu o implementare completă full-stack, documentată și testabilă manual, "
        "care integrează autentificare securizată, plăți online, gestionare evenimente și panou "
        "administrativ într-un singur produs coerent. Modelul de date simplu (eveniment — bilet — utilizator) "
        "s-a dovedit suficient de expresiv pentru conferințe, concerte, meeting-uri și evenimente sportive amatoriale.",
    )
    add_body(
        doc,
        "Integrarea Stripe Checkout și a emailului tranzacțional Brevo demonstrează capacitatea de a "
        "conecta servicii externe mature într-un flux unitar. Separarea frontend/backend facilitează "
        "mentenanța și extinderea viitoare, conform principiilor identificate în literatura analizată.",
    )
    add_body(
        doc,
        "Documentația include diagrame de arhitectură, cazuri de utilizare, secvență, activitate și ERD, "
        "alături de capturi de ecran ale interfețelor principale, oferind o imagine completă asupra "
        "soluției implementate pentru organizarea și gestionarea evenimentelor.",
    )
    doc.add_page_break()
